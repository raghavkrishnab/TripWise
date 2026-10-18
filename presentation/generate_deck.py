from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
SCREENSHOTS = ROOT / "screenshots"
OUTPUT = ROOT / "TripWise-Presentation.pptx"

INK = RGBColor(27, 48, 37)
FOREST = RGBColor(40, 91, 67)
GOLD = RGBColor(187, 135, 80)
PAPER = RGBColor(247, 246, 240)
SAGE = RGBColor(229, 238, 230)
MUTED = RGBColor(105, 119, 108)
WHITE = RGBColor(255, 255, 251)
LINE = RGBColor(218, 222, 211)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]


def set_background(slide, color=PAPER):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def text(slide, x, y, w, h, value, size=18, color=INK, bold=False,
         font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP,
         margin=0.02, italic=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return box


def rect(slide, x, y, w, h, fill, line=None, radius=False):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    if radius:
        shape.adjustments[0] = 0.08
    return shape


def title(slide, eyebrow, heading, subtitle=None, number=None):
    text(slide, 0.72, 0.42, 10.9, 0.25, eyebrow.upper(), 10, GOLD, True)
    text(slide, 0.7, 0.82, 11.9, 0.6, heading, 29, INK, True, "Aptos Display")
    if subtitle:
        text(slide, 0.74, 1.52, 11.4, 0.38, subtitle, 13, MUTED)
    if number:
        text(slide, 12.1, 7.05, 0.45, 0.2, f"{number:02d}", 9, MUTED, align=PP_ALIGN.RIGHT)
    rect(slide, 0.74, 6.94, 0.48, 0.035, GOLD)
    text(slide, 1.32, 6.84, 4.8, 0.22, "TRIPWISE  /  SMART TRAVEL PLANNER", 8, MUTED, True)


def bullet(slide, x, y, w, heading, detail, accent=FOREST):
    rect(slide, x, y + 0.05, 0.08, 0.58, accent, radius=True)
    text(slide, x + 0.22, y, w - 0.22, 0.28, heading, 15, INK, True)
    text(slide, x + 0.22, y + 0.31, w - 0.22, 0.54, detail, 11, MUTED)


def screenshot(slide, name, x, y, w, h, label):
    path = SCREENSHOTS / name
    if not path.exists():
        raise FileNotFoundError(path)
    frame = rect(slide, x, y, w, h, WHITE, LINE, radius=True)
    frame.shadow.inherit = False
    frame.shadow.blur = Pt(8)
    frame.shadow.distance = Pt(2)
    frame.shadow.angle = 45
    frame.shadow.opacity = 0.12
    slide.shapes.add_picture(str(path), Inches(x + 0.08), Inches(y + 0.08), width=Inches(w - 0.16), height=Inches(h - 0.16))
    text(slide, x, y + h + 0.09, w, 0.24, label, 9, MUTED, True, align=PP_ALIGN.CENTER)


# 1. Cover
slide = prs.slides.add_slide(blank)
set_background(slide, INK)
slide.shapes.add_picture(str(SCREENSHOTS / "01-explore.png"), Inches(6.8), Inches(0.0), width=Inches(6.54), height=Inches(7.5))
rect(slide, 6.8, 0, 0.12, 7.5, GOLD)
text(slide, 0.86, 0.72, 4.9, 0.3, "FULL-STACK PROJECT PRESENTATION", 10, RGBColor(226, 194, 143), True)
text(slide, 0.82, 1.7, 5.6, 1.35, "TripWise", 54, WHITE, True, "Aptos Display")
text(slide, 0.88, 3.08, 5.2, 1.15, "Discover places.\nBuild a trip worth taking.", 24, RGBColor(225, 232, 220), False, "Aptos Display")
rect(slide, 0.9, 4.67, 0.75, 0.055, GOLD)
text(slide, 0.88, 5.0, 5.1, 0.58, "React  ·  Express  ·  PostgreSQL  ·  Geoapify", 14, WHITE, True)
text(slide, 0.88, 6.55, 4.5, 0.3, "Single-service deployment plan for Render", 11, RGBColor(195, 207, 195))

# 2. Product overview
slide = prs.slides.add_slide(blank)
set_background(slide)
title(slide, "The product", "One place to go from discovery to itinerary", "TripWise combines real place search with a persistent, shared travel shortlist.", 2)
cards = [
    (0.78, "01", "Search real places", "Geocode a city and retrieve nearby destinations from Geoapify."),
    (4.82, "02", "Refine the results", "Choose a category, sort by name, or narrow the list with text search."),
    (8.86, "03", "Keep the good ones", "Open a destination, save it to My Trip, and remove it when plans change."),
]
for x, num, heading, detail in cards:
    rect(slide, x, 2.35, 3.65, 2.2, WHITE, LINE, radius=True)
    text(slide, x + 0.25, 2.62, 0.6, 0.34, num, 13, GOLD, True)
    text(slide, x + 0.25, 3.08, 3.1, 0.45, heading, 17, INK, True)
    text(slide, x + 0.25, 3.62, 3.08, 0.74, detail, 12, MUTED)
text(slide, 0.85, 5.13, 11.5, 0.55, "Designed for a fast demo: search → inspect → save → revisit.", 20, FOREST, True, "Aptos Display")
text(slide, 0.86, 5.88, 11.2, 0.5, "The deployed MVP has no accounts; the itinerary is shared by all visitors to the same service.", 12, MUTED)

# 3. Explore UI
slide = prs.slides.add_slide(blank)
set_background(slide)
title(slide, "Explore", "Live destinations, useful controls", "Search results come from Geoapify; category and sort state remain in sync.", 3)
screenshot(slide, "01-explore.png", 0.78, 2.02, 7.65, 4.28, "EXPLORE: DESTINATION SEARCH AND RESULTS")
bullet(slide, 8.9, 2.25, 3.55, "City search", "Resolve a place name, then fetch nearby destinations.")
bullet(slide, 8.9, 3.36, 3.55, "Category filters", "Attractions, sights, restaurants, entertainment, and hotels.", GOLD)
bullet(slide, 8.9, 4.47, 3.55, "Find and sort", "Filter the visible list and sort names ascending or descending.")

# 4. Details and itinerary
slide = prs.slides.add_slide(blank)
set_background(slide)
title(slide, "Plan", "A small, complete planning loop", "Destination details lead directly into a database-backed itinerary.", 4)
screenshot(slide, "02-destination-details.png", 0.78, 2.02, 5.7, 3.8, "DESTINATION DETAILS")
screenshot(slide, "04-my-trip-saved.png", 6.86, 2.02, 5.7, 3.8, "MY TRIP: SAVED DESTINATION")
text(slide, 0.93, 6.16, 11.3, 0.38, "Remove confirmation protects against accidental deletes; a reload confirms persistence.", 12, MUTED, align=PP_ALIGN.CENTER)

# 5. Architecture
slide = prs.slides.add_slide(blank)
set_background(slide)
title(slide, "Architecture", "One Node service, one origin", "The Express server serves both the JSON API and Vite's production build.", 5)
nodes = [
    (0.86, 2.6, 2.5, 1.35, "React + Vite", "UI and same-origin fetch"),
    (4.12, 2.6, 2.5, 1.35, "Express", "API, static assets, health"),
    (7.38, 2.6, 2.1, 1.35, "Geoapify", "Geocode + nearby places"),
    (10.22, 2.6, 2.25, 1.35, "PostgreSQL", "Trips + itinerary items"),
]
for x, y, w, h, heading, detail in nodes:
    rect(slide, x, y, w, h, WHITE, LINE, radius=True)
    text(slide, x + 0.17, y + 0.2, w - 0.34, 0.36, heading, 16, FOREST, True, align=PP_ALIGN.CENTER)
    text(slide, x + 0.16, y + 0.68, w - 0.32, 0.42, detail, 10, MUTED, align=PP_ALIGN.CENTER)
for x in (3.45, 6.73, 9.6):
    text(slide, x, 2.96, 0.6, 0.34, "→", 23, GOLD, True, align=PP_ALIGN.CENTER)
rect(slide, 2.1, 4.65, 9.15, 0.82, SAGE, SAGE, radius=True)
text(slide, 2.3, 4.86, 8.75, 0.35, "Render Web Service  ·  npm run build  →  npm start  ·  /api/health", 14, INK, True, align=PP_ALIGN.CENTER)
text(slide, 1.0, 5.92, 11.4, 0.55, "Secrets stay in Render environment settings: DATABASE_URL and GEOAPIFY_API_KEY. No Vercel or client API key.", 12, MUTED, align=PP_ALIGN.CENTER)

# 6. Data flow and reliability
slide = prs.slides.add_slide(blank)
set_background(slide)
title(slide, "Data flow", "Search outside; persist the user's shortlist", "Provider credentials stay server-side; itinerary writes are parameterized SQL.", 6)
bullet(slide, 0.95, 2.18, 5.5, "1. Resolve the city", "Server calls Geoapify geocoding and uses the returned coordinates.")
bullet(slide, 0.95, 3.28, 5.5, "2. Fetch nearby places", "Category, radius, result limit, and requested ordering shape results.", GOLD)
bullet(slide, 0.95, 4.38, 5.5, "3. Save or remove", "Express validates inputs; PostgreSQL enforces one saved place per trip/place ID.")
screenshot(slide, "05-remove-confirmation.png", 7.18, 2.02, 4.8, 3.92, "REMOVE CONFIRMATION")
text(slide, 0.97, 5.92, 5.55, 0.48, "Health check reports DB state, uptime, and application version.", 12, FOREST, True)

# 7. Verification
slide = prs.slides.add_slide(blank)
set_background(slide)
title(slide, "Verified locally", "Release checks that were actually run", "Production-style server served the built frontend and API from the same origin.", 7)
checks = [
    ("Build", "Root Render build completed successfully."),
    ("Search", "Live Geoapify search returned 30 places for tested cities/categories."),
    ("Persistence", "Add, read-back, reload, remove, and reload were verified in PostgreSQL."),
    ("Runtime", "Health returned connected DB + version; static assets and favicon returned 200."),
    ("UX", "Theme persistence, category/sort, confirmation, and 390px no-overflow checks passed."),
]
for index, (heading, detail) in enumerate(checks):
    y = 2.06 + index * 0.78
    rect(slide, 0.95, y + 0.02, 0.31, 0.31, SAGE, SAGE, radius=True)
    text(slide, 0.98, y + 0.035, 0.25, 0.25, "✓", 14, FOREST, True, align=PP_ALIGN.CENTER)
    text(slide, 1.52, y, 1.55, 0.3, heading, 14, INK, True)
    text(slide, 3.12, y, 8.8, 0.45, detail, 12, MUTED)
rect(slide, 0.95, 6.15, 11.3, 0.45, RGBColor(244, 235, 218), RGBColor(244, 235, 218), radius=True)
text(slide, 1.12, 6.25, 10.9, 0.22, "Scope note: no automated test suite is currently included; manual integration checks used live services.", 10, INK)

# 8. Deployment and close
slide = prs.slides.add_slide(blank)
set_background(slide, INK)
text(slide, 0.84, 0.72, 4.6, 0.28, "DEPLOYMENT", 10, RGBColor(226, 194, 143), True)
text(slide, 0.8, 1.18, 11.8, 0.72, "Ready for Render. Not yet publicly deployed.", 30, WHITE, True, "Aptos Display")
rect(slide, 0.9, 2.28, 0.72, 0.05, GOLD)
steps = [
    (0.88, "01", "Push to GitHub", "Use the TripWise repository root only."),
    (4.72, "02", "Create Render service", "Use render.yaml Blueprint or manual Web Service."),
    (8.56, "03", "Set secrets + deploy", "DATABASE_URL, GEOAPIFY_API_KEY; Render sets PORT."),
]
for x, number, heading, detail in steps:
    text(slide, x, 2.85, 0.65, 0.35, number, 13, RGBColor(226, 194, 143), True)
    text(slide, x, 3.35, 3.35, 0.4, heading, 16, WHITE, True)
    text(slide, x, 3.88, 3.25, 0.75, detail, 11, RGBColor(195, 207, 195))
text(slide, 0.9, 5.55, 11.1, 0.5, "Thank you", 27, WHITE, True, "Aptos Display")
text(slide, 0.93, 6.22, 10.9, 0.36, "TripWise  ·  Discover places. Keep the ones that matter.", 13, RGBColor(226, 194, 143))
text(slide, 12.05, 7.05, 0.45, 0.2, "08", 9, RGBColor(195, 207, 195), align=PP_ALIGN.RIGHT)

prs.core_properties.title = "TripWise — Smart Travel Planner"
prs.core_properties.subject = "Full-stack travel planner and Render deployment plan"
prs.core_properties.author = "TripWise project team"
prs.core_properties.keywords = "TripWise, React, Express, PostgreSQL, Geoapify, Render"
prs.save(OUTPUT)
print(f"Created {OUTPUT}")